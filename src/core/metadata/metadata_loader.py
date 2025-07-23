"""
Flexible loader for various database formats into DuckDB.
"""
import time
import pandas as pd
import duckdb
from pathlib import Path
from typing import Dict, List, Any, Optional, Union
import json

from .schemas.base_schema import BaseSchema
from .database_registry import database_registry
from ...config.logging import get_logger
from ...config.settings import settings

logger = get_logger(__name__)

class MetadataLoader:
    """Loads data from various sources into DuckDB for fast querying."""
    
    def __init__(self, db_path: Optional[str] = None):
        # Use in-memory database by default for Railway compatibility
        self.db_path = db_path or ":memory:"
        self.connection = duckdb.connect(self.db_path)
        self._initialized_tables = set()
        
    def load_database(self, 
                     name: str,
                     source_path: str,
                     schema: BaseSchema,
                     force_reload: bool = False) -> int:
        """Load a database from source file into DuckDB."""
        start_time = time.time()
        
        # Check if table already exists and skip if not forcing reload
        if not force_reload and self._table_exists(schema.table_name):
            logger.info(f"Table '{schema.table_name}' already loaded, skipping")
            return self._get_row_count(schema.table_name)
        
        # Load source data
        source_path = Path(source_path)
        if not source_path.exists():
            logger.error(f"Source file not found: {source_path}")
            return 0
        
        # Load based on file type
        if source_path.suffix in ['.xlsx', '.xls']:
            data = self._load_excel(source_path)
        elif source_path.suffix == '.csv':
            data = self._load_csv(source_path)
        elif source_path.suffix == '.json':
            data = self._load_json(source_path)
        else:
            logger.error(f"Unsupported file format: {source_path.suffix}")
            return 0
        
        if not data:
            logger.error(f"No data loaded from {source_path}")
            return 0
        
        # Create table
        self._create_table(schema)
        
        # Transform and insert data
        row_count = self._insert_data(schema, data)
        
        # Update registry
        database_registry.update_database_stats(name, row_count, time.time())
        
        elapsed = time.time() - start_time
        logger.info(f"Loaded {row_count} rows into '{schema.table_name}' in {elapsed:.2f}s")
        
        return row_count
    
    def _load_excel(self, path: Path) -> List[Dict[str, Any]]:
        """Load data from Excel file."""
        try:
            # Check if openpyxl is available
            try:
                import openpyxl
            except ImportError:
                logger.error("openpyxl package not installed. Please install with: pip install openpyxl")
                raise ImportError("openpyxl is required to read Excel files")
            
            # Try to find the main data sheet
            excel_file = pd.ExcelFile(path)
            
            # Look for sheets with data
            data_sheet_names = ['AI Risk Database v3', 'AI Risk Database', 'Data', 'Risks']
            sheet_name = None
            
            for name in data_sheet_names:
                if name in excel_file.sheet_names:
                    sheet_name = name
                    break
            
            if not sheet_name and excel_file.sheet_names:
                # Use first sheet if no match found
                sheet_name = excel_file.sheet_names[0]
                logger.warning(f"Using first sheet '{sheet_name}' as no standard sheet name found")
            
            # Read the sheet
            df = pd.read_excel(path, sheet_name=sheet_name)
            
            # Convert to list of dicts
            data = df.to_dict('records')
            logger.info(f"Loaded {len(data)} rows from Excel sheet '{sheet_name}'")
            
            return data
            
        except ImportError as e:
            logger.error(f"Import error: {str(e)}")
            raise
        except Exception as e:
            logger.error(f"Error loading Excel file: {str(e)}")
            return []
    
    def _load_csv(self, path: Path) -> List[Dict[str, Any]]:
        """Load data from CSV file."""
        try:
            df = pd.read_csv(path)
            return df.to_dict('records')
        except Exception as e:
            logger.error(f"Error loading CSV file: {str(e)}")
            return []
    
    def _load_json(self, path: Path) -> List[Dict[str, Any]]:
        """Load data from JSON file."""
        try:
            with open(path, 'r') as f:
                data = json.load(f)
                if isinstance(data, list):
                    return data
                elif isinstance(data, dict) and 'data' in data:
                    return data['data']
                else:
                    return [data]
        except Exception as e:
            logger.error(f"Error loading JSON file: {str(e)}")
            return []
    
    def _create_table(self, schema: BaseSchema):
        """Create table in DuckDB based on schema."""
        try:
            # Drop table if exists for reload
            self.connection.execute(f"DROP TABLE IF EXISTS {schema.table_name}")
            
            # Create table
            create_sql = schema.get_sql_create_table()
            self.connection.execute(create_sql)
            
            # Create indices
            for index_sql in schema.get_sql_indices():
                self.connection.execute(index_sql)
            
            self._initialized_tables.add(schema.table_name)
            logger.info(f"Created table '{schema.table_name}' with indices")
            
        except Exception as e:
            logger.error(f"Error creating table: {str(e)}")
            raise
    
    def _insert_data(self, schema: BaseSchema, data: List[Dict[str, Any]]) -> int:
        """Insert data into table using schema transformation."""
        inserted = 0
        batch_size = 1000
        
        # Process in batches for better performance
        for i in range(0, len(data), batch_size):
            batch = data[i:i + batch_size]
            batch_data = []
            
            for row in batch:
                try:
                    # Validate row
                    if not schema.validate_row(row):
                        continue
                    
                    # Transform row using schema
                    transformed = schema.extract_row_data(row)
                    batch_data.append(transformed)
                    
                except Exception as e:
                    logger.warning(f"Error processing row: {str(e)}")
                    continue
            
            if batch_data:
                # Insert batch
                try:
                    columns = list(batch_data[0].keys())
                    placeholders = ', '.join(['?' for _ in columns])
                    insert_sql = f"INSERT INTO {schema.table_name} ({', '.join(columns)}) VALUES ({placeholders})"
                    
                    # Convert to tuples for insertion
                    values = [tuple(row[col] for col in columns) for row in batch_data]
                    self.connection.executemany(insert_sql, values)
                    
                    inserted += len(batch_data)
                    
                except Exception as e:
                    logger.error(f"Error inserting batch: {str(e)}")
        
        return inserted
    
    def _table_exists(self, table_name: str) -> bool:
        """Check if a table exists in the database."""
        try:
            result = self.connection.execute(
                "SELECT COUNT(*) FROM information_schema.tables WHERE table_name = ?",
                [table_name]
            ).fetchone()
            return result[0] > 0
        except:
            return False
    
    def _get_row_count(self, table_name: str) -> int:
        """Get the number of rows in a table."""
        try:
            result = self.connection.execute(f"SELECT COUNT(*) FROM {table_name}").fetchone()
            return result[0]
        except:
            return 0
    
    def execute_query(self, sql: str, params: Optional[List[Any]] = None) -> pd.DataFrame:
        """Execute a SQL query and return results as DataFrame."""
        try:
            if params:
                result = self.connection.execute(sql, params)
            else:
                result = self.connection.execute(sql)
            
            return result.df()
            
        except Exception as e:
            logger.error(f"Error executing query: {str(e)}")
            raise
    
    def get_connection(self) -> duckdb.DuckDBPyConnection:
        """Get the DuckDB connection for direct queries."""
        return self.connection
    
    def close(self):
        """Close the database connection."""
        if self.connection:
            self.connection.close()